#!/usr/bin/env python3
"""Synthetic parser fixtures (requires python-docx, reportlab, Pillow).

XLSX is authored separately with Artifact Tool. Legacy DOC/XLS are converted
with LibreOffice; see data/README.md. No real invoices or personal information.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.shared import RGBColor
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/attachments'
for d in ['office','pdf','images','invalid']:
    (OUT / d).mkdir(parents=True, exist_ok=True)
rows=json.loads((ROOT/'data/expected/attachment-rows.json').read_text())['rows']
font_paths = ['/System/Library/Fonts/Supplemental/Arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
font_path=next(p for p in font_paths if Path(p).exists())
pdfmetrics.registerFont(TTFont('Fixture',font_path))

doc=Document()
doc.sections[0].left_margin=Inches(.7)
doc.sections[0].right_margin=Inches(.7)
title=doc.add_heading('Синтетическая спецификация',0)
for run in title.runs: run.font.color.rgb=RGBColor(0,0,0)
for element in [title._p, doc.styles['Title'].element]:
    for border in list(element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.add_paragraph('SYNTHETIC FIXTURE. Данные для тестирования, не заказ.')
table=doc.add_table(rows=1,cols=5)
table.style='Light Shading Accent 1'
for c,t in zip(table.rows[0].cells,['Артикул','Описание','Количество','Единица','Цена']): c.text=t
for i,r in enumerate(rows):
    for c,t in zip(table.add_row().cells,[r['article'],r['description'],(r['quantity'] or '').replace('.',','),r['unit'] or '', '1500' if i==0 else '2700' if i==1 else '']): c.text=t
for p in doc.paragraphs:
    for run in p.runs: run.font.size=Pt(11 if p.style.name=='Normal' else 18)
doc.save(OUT/'office/specification.docx')

lines=['SYNTHETIC FIXTURE','Артикул | Описание | Количество | Единица | Цена']+[
    ' | '.join([r['article'],r['description'],(r['quantity'] or '').replace('.',','),r['unit'] or '', '1500' if i==0 else '2700' if i==1 else ''])
    for i,r in enumerate(rows)]
textpdf=canvas.Canvas(str(OUT/'pdf/text.pdf'),pagesize=(842,595))
textpdf.setFont('Fixture',13)
for i,line in enumerate(lines): textpdf.drawString(35,550-i*35,line)
textpdf.save()
font=ImageFont.truetype(font_path,26)
scan=Image.new('RGB',(1684,1190),'white')
d=ImageDraw.Draw(scan)
for i,line in enumerate(lines): d.text((70,80+i*70),line,font=font,fill='black')
scan.save(OUT/'images/label.jpg',quality=95)
scan.rotate(90,expand=True).save(OUT/'images/rotated.jpg',quality=90)
scan.filter(ImageFilter.GaussianBlur(12)).save(OUT/'images/blurred.jpg',quality=65)
pdf=canvas.Canvas(str(OUT/'pdf/scan.pdf'),pagesize=(842,595))
pdf.drawImage(ImageReader(scan),0,0,width=842,height=595)
pdf.save()
pdf=canvas.Canvas(str(OUT/'pdf/mixed.pdf'),pagesize=(842,595))
pdf.setFont('Fixture',13)
for i,line in enumerate(lines): pdf.drawString(35,550-i*35,line)
pdf.showPage()
pdf.drawImage(ImageReader(scan),0,0,width=842,height=595)
pdf.save()
# Deliberately unmarked synthetic product shape: no hidden ratings may be inferred.
product=Image.new('RGB',(400,600),'#eeeeee')
d=ImageDraw.Draw(product)
d.rounded_rectangle((110,50,290,550),15,fill='#dddddd',outline='#555555',width=5)
d.rectangle((140,180,260,350),fill='#333333')
d.rectangle((155,200,245,255),fill='#ee6622')
d.ellipse((172,72,228,128),fill='#777777')
d.ellipse((172,472,228,528),fill='#777777')
product.save(OUT/'images/product-only.jpg',quality=95)
(OUT/'invalid/corrupt.pdf').write_bytes(b'%PDF-1.7\ninvalid synthetic fixture')
(OUT/'invalid/renamed-executable.jpg').write_bytes(b'MZ SYNTHETIC NOT AN IMAGE')
certdir=ROOT/'data/sample_catalog/certificates'
certdir.mkdir(exist_ok=True)
pdf=canvas.Canvas(str(certdir/'synthetic-certificate.pdf'))
pdf.setFont('Fixture',14)
pdf.drawString(40,780,'SYNTHETIC CERTIFICATE - NOT VALID FOR SALE')
pdf.drawString(40,750,'Тестовые категории: breakers, cables, lamps')
pdf.drawString(40,720,'Не является подтверждением электробезопасности.')
pdf.save()
print('Generated Office/PDF/JPEG fixtures. Convert legacy files, then run scripts/update-fixture-manifest.py.')
